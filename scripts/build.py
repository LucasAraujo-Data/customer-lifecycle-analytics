"""Phase 2 ingestion and point-in-time feature build. Run from repository root."""
import argparse
import hashlib
import json
from datetime import datetime, timedelta
from pathlib import Path
import duckdb
from openpyxl import load_workbook

HEADERS = ['Invoice','StockCode','Description','Quantity','InvoiceDate','Price','Customer ID','Country']

def ingest(con, source, out, batch_size=10000):
    con.execute('''CREATE TABLE raw_lines (source_sheet VARCHAR, source_row_number BIGINT, invoice VARCHAR,
       stock_code VARCHAR, description VARCHAR, quantity DOUBLE, invoice_timestamp TIMESTAMP,
       price DOUBLE, customer_id VARCHAR, country VARCHAR)''')
    wb = load_workbook(source, read_only=True, data_only=True)
    counts = {}
    for ws in wb:
        rows = ws.iter_rows(values_only=True)
        headers = list(next(rows))[:8]
        if headers != HEADERS:
            raise ValueError(f'Unexpected headers in {ws.title}: {headers}')
        buffer=[]; count=0
        for row_number, row in enumerate(rows, start=2):
            vals=list(row[:8]); invoice, stock, desc, qty, date, price, cid, country = vals
            if isinstance(date, datetime): stamp=date
            elif date is None: stamp=None
            else: raise ValueError(f'Unexpected date at {ws.title}:{row_number}: {date!r}')
            cid = None if cid is None else str(int(cid)) if isinstance(cid,(float,int)) and float(cid).is_integer() else str(cid).strip()
            buffer.append((ws.title,row_number,str(invoice).strip() if invoice is not None else None,
                str(stock).strip() if stock is not None else None,desc,qty,stamp,price,cid,country))
            count+=1
            if len(buffer)>=batch_size:
                con.executemany('INSERT INTO raw_lines VALUES (?,?,?,?,?,?,?,?,?,?)',buffer); buffer=[]
        if buffer: con.executemany('INSERT INTO raw_lines VALUES (?,?,?,?,?,?,?,?,?,?)',buffer)
        counts[ws.title]=count
    wb.close()
    if counts != {'Year 2009-2010':525461,'Year 2010-2011':541910}:
        raise AssertionError(f'Unexpected source sheet counts: {counts}')
    con.execute(f"COPY raw_lines TO '{(out/'raw_lines.parquet').as_posix()}' (FORMAT PARQUET)")
    return counts

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--source',required=True,type=Path)
    ap.add_argument('--output',default='output',type=Path)
    args=ap.parse_args(); source=args.source.resolve(); out=args.output.resolve(); out.mkdir(parents=True,exist_ok=True)
    if not source.is_file(): raise FileNotFoundError(source)
    checksum=hashlib.sha256()
    with source.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''): checksum.update(chunk)
    con=duckdb.connect(str(out/'phase2.duckdb'))
    try:
        counts=ingest(con,source,out)
        sql_path=Path(__file__).resolve().parent.parent/'sql'/'transform.sql'
        con.execute(sql_path.read_text(encoding='utf-8'))
        checks=Path(__file__).resolve().parent.parent/'tests'/'checks.sql'
        con.execute(checks.read_text(encoding='utf-8'))
        metrics={r[0]:r[1] for r in con.execute('SELECT metric, value FROM quality_metrics').fetchall()}
        for name in ('quality_events','orders_clean','customer_snapshot','customer_features'):
            con.execute(f"COPY {name} TO '{(out/(name+'.parquet')).as_posix()}' (FORMAT PARQUET)")
        manifest={'source_sha256':checksum.hexdigest(),'source_name':source.name,'sheet_rows':counts,
                  'run_timestamp_utc':datetime.utcnow().isoformat()+'Z','quality_metrics':metrics,
                  'feature_version':'v1','status':'completed'}
        (out/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
        print(json.dumps(manifest,indent=2))
    finally: con.close()
if __name__=='__main__':main()
